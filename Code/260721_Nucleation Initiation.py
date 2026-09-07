# -*- coding: utf-8 -*-
"""
PL Signal Detection & Region Timeline Suite (Startup / Time Before PL Detected)
=============================================================================
1. Reads Region 1 start bounds from .docx files.
2. Plots two versions: one with black outlines and one without outlines.
"""

import os
import re
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.transforms import Bbox
from docx import Document

# Global typography and aesthetic setup
plt.rcParams['font.family'] = 'sans-serif'
plt.rcParams['font.sans-serif'] = ['Arial', 'Helvetica', 'DejaVu Sans']
plt.rcParams['savefig.transparent'] = True

# =============================================================================
# DATA PIPELINE CONFIGURATION
# =============================================================================
PL_SAMPLES = [
    {
        "label": "Control",
        "docx_path": r"C:\Users\matth\OneDrive\Desktop\Manuscript data\260717_Data\Reports\260717_PL_FitResults_control.docx",
        "color": "#000000"  # Black
    },
    {
        "label": "3-APA",
        "docx_path": r"C:\Users\matth\OneDrive\Desktop\Manuscript data\260717_Data\Reports\260717_PL_FitResults_3APA.docx",
        "color": "#E07A5F"  # 1st Hex (Coral / Terracotta)
    },
    {
        "label": "4-ABA",
        "docx_path": r"C:\Users\matth\OneDrive\Desktop\Manuscript data\260717_Data\Reports\260717_PL_FitResults_4ABA.docx",
        "color": "#38B000"  # 2nd Hex (Kelly Green)
    },
    {
        "label": "5-AVA",
        "docx_path": r"C:\Users\matth\OneDrive\Desktop\Manuscript data\260717_Data\Reports\260717_PL_FitResults_5AVA.docx",
        "color": "#A381C6"  # 3rd Hex (Muted Purple)
    },
    {
        "label": "7-AHA",
        "docx_path": r"C:\Users\matth\OneDrive\Desktop\Manuscript data\260717_Data\Reports\260717_PL_FitResults_7AHA.docx",
        "color": "#70E4BC"  # 4th Hex (Mint / Soft Teal)
    }
]

# =============================================================================
# WORD DOCUMENT PARSER ENGINE
# =============================================================================

def parse_docx_region_bounds(docx_path):
    """
    Parses a Word document (.docx) to extract Region 1 start timestamps.
    """
    doc = Document(docx_path)

    all_lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            all_lines.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    all_lines.append(cell.text.strip())

    full_text = "\n".join(all_lines)

    def find_region_range(region_num):
        pattern = rf"Region\s*{region_num}[:\s|-]*([\d\.]+)\s*to\s*([\d\.]+)"
        match = re.search(pattern, full_text, re.IGNORECASE)
        if match:
            return float(match.group(1)), float(match.group(2))
        return 0.0, 0.0

    r1_start, _ = find_region_range(1)

    return {
        'r1_start': r1_start
    }


def batch_process_docx_files():
    """Loops through sample configs and reads Region 1 start bounds from .docx files."""
    results = []

    for cfg in PL_SAMPLES:
        try:
            bounds = parse_docx_region_bounds(cfg['docx_path'])
            bounds['label'] = cfg['label']
            bounds['color'] = cfg['color']
            results.append(bounds)

            print(f"✓ Parsed {cfg['label']}:")
            print(f"   Time Before PL Detected: 0.0s -> {bounds['r1_start']}s")
        except Exception as e:
            print(f"[!] Could not parse .docx for {cfg['label']} ({cfg['docx_path']}): {e}")

    return pd.DataFrame(results)

# =============================================================================
# PLOTTING FUNCTIONS
# =============================================================================

def plot_region_stacked_chart(df_results, with_outline=True, save_filename="pl_startup_timeline_square.png"):
    """
    Generates an exact 1:1 pixel square chart.
    - with_outline=True: Adds a black border around each bar.
    - with_outline=False: No borders around bars.
    """
    fig = plt.figure(figsize=(6.0, 6.0), dpi=300)
    fig.patch.set_alpha(0.0)

    ax = fig.add_axes([0.13, 0.13, 0.82, 0.82])
    ax.set_facecolor('none')
    ax.set_box_aspect(1.0)

    labels = df_results['label']
    colors = df_results['color']
    width = 0.55
    x_coords = np.arange(len(labels))

    heights = df_results['r1_start']
    bottoms = np.zeros_like(heights)

    max_val = heights.max()
    dynamic_limit = max_val + 5.0

    edge_color = 'black' if with_outline else 'none'
    line_width = 1.2 if with_outline else 0.0

    for i in range(len(x_coords)):
        ax.bar(
            x_coords[i],
            heights[i],
            width,
            bottom=bottoms[i],
            color=colors[i],
            edgecolor=edge_color,
            linewidth=line_width,
            label=labels[i]
        )

    # Axis Setup & Formatting
    ax.set_xticks(x_coords)
    ax.set_xticklabels(labels)

    ax.set_ylabel('Time before PL Detected (s)', fontsize=14, fontname='Arial', fontweight='normal', labelpad=6)
    ax.set_ylim(0, dynamic_limit)

    ax.tick_params(axis='x', bottom=False, top=False)
    ax.tick_params(axis='y', direction='in', right=True, width=0.8, length=5, labelsize=14)

    for tick in ax.get_xticklabels():
        tick.set_fontname('Arial')
        tick.set_fontsize(14)
        tick.set_fontweight('normal')

    for tick in ax.get_yticklabels():
        tick.set_fontname('Arial')
        tick.set_fontsize(14)
        tick.set_fontweight('normal')

    for spine in ax.spines.values():
        spine.set_linewidth(0.8)

    ax.legend(
        loc='upper left',
        frameon=False,
        handlelength=1.2,
        handletextpad=0.4,
        labelspacing=0.4,
        borderaxespad=0.5,
        prop={'family': 'Arial', 'size': 12, 'weight': 'normal'}
    )

    # Measure tight bounding box
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    tight_bbox = fig.get_tightbbox(renderer)

    pad = 0.04
    tight_bbox = tight_bbox.padded(pad)

    content_w, content_h = tight_bbox.width, tight_bbox.height
    side = max(content_w, content_h)

    x0 = tight_bbox.x0 - (side - content_w) / 2.0
    y0 = tight_bbox.y0 - (side - content_h) / 2.0
    square_bbox = Bbox.from_bounds(x0, y0, side, side)

    save_path = os.path.join(r"C:\Users\matth\Downloads", save_filename)
    plt.savefig(
        save_path,
        dpi=300,
        transparent=True,
        bbox_inches=square_bbox
    )
    plt.show()

# =============================================================================
# MAIN RUNNER
# =============================================================================
if __name__ == "__main__":
    df_detections = batch_process_docx_files()

    if not df_detections.empty:
        # Option 1: Plot WITH black outlines
        print("\n--- Generating Chart WITH Outlines ---")
        plot_region_stacked_chart(
            df_detections,
            with_outline=True,
            save_filename="pl_startup_timeline_square_outlined.png"
        )

        # Option 2: Plot WITHOUT black outlines
        print("\n--- Generating Chart WITHOUT Outlines ---")
        plot_region_stacked_chart(
            df_detections,
            with_outline=False,
            save_filename="pl_startup_timeline_square_no_outline.png"
        )
